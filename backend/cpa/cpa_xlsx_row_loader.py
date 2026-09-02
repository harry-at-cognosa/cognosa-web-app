"""
Row-oriented loader for Ann_List_of_Sources-200704f.xlsx into cpa-c42t_collection.

This spreadsheet is a table ABOUT documents - one row per source - not a document.
UnstructuredExcelLoader would flatten it to prose and destroy the row structure, so
it gets its own path. See PLAN.md section 8.

Point IDs are uuid5 over (row_key, chunk_index) rather than over the chunk text as
in cpa_document_loader_qdrant.py. Content hashing is right for prose, where identical
text is genuinely redundant; here two rows could legitimately share text and must stay
distinct. Keying on the row makes the loader idempotent without collapsing rows.
"""
import argparse, re, unicodedata, uuid
from pathlib import Path

import openpyxl
from langchain_core.documents import Document
from langchain_huggingface import HuggingFaceEmbeddings
from langchain_qdrant import QdrantVectorStore
from langchain_text_splitters import RecursiveCharacterTextSplitter
from qdrant_client import QdrantClient

MODEL         = "sentence-transformers/all-MiniLM-L6-v2"
CHUNK_SIZE    = 1000
CHUNK_OVERLAP = 200
FIRST_ROW, LAST_ROW = 2, 348
CPA_NAMESPACE = uuid.UUID("6f2a1c94-3b7d-5e18-9a4f-c42b00000000")

# Columns A-O. Column P is blank apart from stray formatting and is ignored.
# Note the header letters do NOT match the sheet columns: header "A - AES How"
# lives in column G. Flag keys follow the HEADER letter, matching the semantic
# scheme rather than sheet position.
COLS = {
    1: 'seqn', 2: 'nseqn', 3: 'citation', 4: 'y', 5: 'z', 6: 'citno',
    7: 'flag_a', 8: 'flag_b', 9: 'flag_c', 10: 'flag_d', 11: 'flag_e', 12: 'flag_f',
    13: 'more_notes', 14: 'key', 15: 'note3',
}
FLAG_COLS = range(7, 13)                      # G-L: metadata only, never embedded
TEXT_COLS = [1, 2, 3, 4, 5, 6, 13, 14, 15]    # what goes into the vector


def as_text(v) -> str:
    """Every cell becomes a string. Integral floats must not render as '9.0'."""
    if v is None:
        return ''
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v).strip()


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('--xlsx', default='/Users/harry/Downloads/3-lit-review-sources-docs/Ann_List_of_Sources-200704f.xlsx')
    ap.add_argument('--sheet', default='END2020_A')
    ap.add_argument('--collection', default='cpa-c42t_collection')
    ap.add_argument('--host', default='localhost')
    ap.add_argument('--port', type=int, default=6333)
    ap.add_argument('--dry-run', action='store_true')
    args = ap.parse_args()

    ws = openpyxl.load_workbook(args.xlsx, data_only=True)[args.sheet]
    headers = {c: as_text(ws.cell(1, c).value) for c in COLS}
    splitter = RecursiveCharacterTextSplitter(
        chunk_size=CHUNK_SIZE, chunk_overlap=CHUNK_OVERLAP, length_function=len)

    docs, ids = [], []
    multi = 0
    for n, r in enumerate(range(FIRST_ROW, LAST_ROW + 1), start=1):
        vals = {c: as_text(ws.cell(r, c).value) for c in COLS}
        if not any(vals.values()):
            continue

        citation = vals[3]
        row_key = f"{n:04d}|{citation[:200]}"

        flags_set   = [COLS[c][-1] for c in FLAG_COLS if vals[c]]
        flag_labels = [f"{headers[c]}: {vals[c]}" for c in FLAG_COLS if vals[c]]

        body = "\n".join(f"{headers[c]}: {vals[c]}" for c in TEXT_COLS if vals[c])

        meta = {COLS[c]: vals[c] for c in COLS}
        meta.update({
            'row_number': n,
            'row_key': row_key,
            'loader_row': r,
            'flags_set': flags_set,
            'flag_labels': flag_labels,
            'source': str(Path(args.xlsx)),
            'sheet': args.sheet,
            'cpa_source_file': Path(args.xlsx).name,
            'cpa_corpus': 'xlsx-rows',
        })

        pieces = splitter.split_text(body) or ['']
        if len(pieces) > 1:
            multi += 1
        for i, piece in enumerate(pieces):
            m = dict(meta, chunk_index=i, chunk_count=len(pieces))
            docs.append(Document(page_content=piece, metadata=m))
            ids.append(str(uuid.uuid5(CPA_NAMESPACE, f"{row_key}#{i}")))

    print(f"rows {FIRST_ROW}-{LAST_ROW}: {len(set(d.metadata['row_number'] for d in docs))} rows "
          f"-> {len(docs)} points ({multi} rows split into multiple)")
    if args.dry_run:
        d = docs[0]
        print("\n--- first point ---")
        print(d.page_content[:300])
        print("metadata:", {k: v for k, v in d.metadata.items() if k not in ('source',)})
        return

    client = QdrantClient(host=args.host, port=args.port)
    store = QdrantVectorStore(client=client, collection_name=args.collection,
                              embedding=HuggingFaceEmbeddings(model_name=MODEL))
    store.add_documents(docs, ids=ids)
    info = client.get_collection(args.collection)
    print(f"upserted. collection points={info.points_count} indexed={info.indexed_vectors_count}")


if __name__ == '__main__':
    main()
